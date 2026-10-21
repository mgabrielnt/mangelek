#!/usr/bin/env python3
"""Pecah BRD/FRD xlsx menjadi Markdown kecil yang hemat token untuk agent.

Jalankan ulang setiap kali BRD/FRD berubah (setelah disetujui Gab):
    python scripts/agents/export-requirements.py
Hasil di docs/requirements/ dan docs/database/erd/ adalah turunan; jangan diedit manual.
"""
import re, sys
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parents[2]
FRD = ROOT / "FRD-Mangelek-v3.0.xlsx"
BRD = ROOT / "BRD-Mangelek-v4.0.xlsx"
OUT_REQ = ROOT / "docs" / "requirements"
OUT_ERD = ROOT / "docs" / "database" / "erd"
NOTE = "<!-- Dibuat otomatis oleh scripts/agents/export-requirements.py dari {src}. Jangan edit manual. -->\n\n"


def cell(v):
    if v is None:
        return ""
    s = str(v).strip().replace("|", "\\|")
    return re.sub(r"\s*\n\s*", "<br>", s)


def slug(s):
    s = s.lower()
    s = re.sub(r"[^a-z0-9]+", "-", s)
    return s.strip("-")[:60]


def rows_of(wb, name):
    return [tuple(r) for r in wb[name].iter_rows(values_only=True)]


def table(header, rows):
    out = ["| " + " | ".join(cell(h) for h in header) + " |",
           "| " + " | ".join("---" for _ in header) + " |"]
    for r in rows:
        out.append("| " + " | ".join(cell(c) for c in r) + " |")
    return "\n".join(out) + "\n"


def sheet_to_md(rows):
    """Baris 1 sel = paragraf; baris >=2 sel = tabel (baris pertama grup = header)."""
    out, group = [], []

    def flush():
        if group:
            width = max(len(g) for g in group)
            g2 = [list(g) + [""] * (width - len(g)) for g in group]
            out.append(table(g2[0], g2[1:]))
            group.clear()

    for r in rows:
        vals = list(r)
        while vals and (vals[-1] is None or str(vals[-1]).strip() == ""):
            vals.pop()
        filled = [v for v in vals if v is not None and str(v).strip() != ""]
        if not filled:
            flush()
            continue
        if len(filled) == 1 and len(vals) <= 1 or (len(filled) == 1 and vals[0] is not None):
            flush()
            out.append(cell(filled[0]).replace("<br>", "\n") + "\n")
        else:
            group.append(vals)
    flush()
    return "\n".join(out)


def write(path, text, src):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(NOTE.format(src=src) + text, encoding="utf-8", newline="\n")
    return path


def header_index(rows, first):
    for i, r in enumerate(rows):
        if r and r[0] == first:
            return i
    raise SystemExit(f"Header '{first}' tidak ditemukan")


def main():
    written = []
    frd = openpyxl.load_workbook(FRD, read_only=True, data_only=True)
    brd = openpyxl.load_workbook(BRD, read_only=True, data_only=True)

    # --- Modul FRD ---
    mod_rows = rows_of(frd, "4.1 Modul")
    hi = header_index(mod_rows, "Kode")
    modules = [r for r in mod_rows[hi + 1:] if r[0] and r[0] != "Total"]
    fr = rows_of(frd, "4.2 Kebutuhan Fungsional"); fh = header_index(fr, "Modul")
    ac = rows_of(frd, "10.2 Kriteria Penerimaan"); ah = header_index(ac, "Modul")
    tc = rows_of(frd, "Katalog Test Case"); th = header_index(tc, "ID")
    tc_head = tc[th][:12]
    tc_rows = [r[:12] for r in tc[th + 1:] if r[0]]
    used_tc = set()
    index_mod = []
    for m in modules:
        code, name, feats = m[0], m[1], m[2]
        frs = [r[1:12] for r in fr[fh + 1:] if r[0] == code]
        acs = [r[1:6] for r in ac[ah + 1:] if r[0] == code]
        tcs = [r for r in tc_rows if r[1] and str(r[1]).startswith(code + " ")]
        used_tc.update(r[0] for r in tcs)
        must = sum(1 for r in frs if r[6] == "Must")
        md = f"# {code} — {name}\n\nFitur utama: {cell(feats)}\n\n" \
             f"Ringkas: {len(frs)} kebutuhan ({must} Must), {len(acs)} kriteria penerimaan, {len(tcs)} test case.\n\n"
        md += "## Kebutuhan fungsional\n\n" + table(fr[fh][1:12], frs) + "\n"
        md += "## Kriteria penerimaan\n\n" + table(ac[ah][1:6], acs) + "\n"
        md += "## Test case\n\n" + table(tc_head, tcs)
        p = write(OUT_REQ / "frd" / "modules" / f"{code}.md", md, FRD.name)
        written.append(p); index_mod.append((code, name, len(frs), must))

    rest = [r for r in tc_rows if r[0] not in used_tc]
    written.append(write(OUT_REQ / "frd" / "test-cases-lintas-modul.md",
                         "# Test case lintas modul (NFR, migrasi, isolasi, pembayaran, E2E, error)\n\n" + table(tc_head, rest), FRD.name))

    # --- Entitas per kelompok (ERD) ---
    en = rows_of(frd, "7.2 Entitas"); eh = header_index(en, "Kelompok")
    groups = {}
    for r in en[eh + 1:]:
        if r[0] and r[1]:
            groups.setdefault(r[0], []).append(r[1:5])
    for g, items in groups.items():
        letter = g.split(".")[0].strip()
        md = f"# Entitas {g}\n\nDiagram: lihat Mangelek-Diagram/mmd/17*-erd-*.mmd. Tanda † = Fase 2.\n\n" + table(en[eh][1:5], items)
        written.append(write(OUT_ERD / f"{letter}-{slug(g.split('.',1)[1])}.md", md, FRD.name))

    # --- Sheet FRD lain ---
    skip = {"4.1 Modul", "4.2 Kebutuhan Fungsional", "10.2 Kriteria Penerimaan", "Katalog Test Case", "7.2 Entitas", "Diagram"}
    for ws in frd.worksheets:
        if ws.title in skip:
            continue
        written.append(write(OUT_REQ / "frd" / f"{slug(ws.title)}.md", sheet_to_md(rows_of(frd, ws.title)), FRD.name))
    for ws in brd.worksheets:
        if ws.title == "Diagram":
            continue
        written.append(write(OUT_REQ / "brd" / f"{slug(ws.title)}.md", sheet_to_md(rows_of(brd, ws.title)), BRD.name))

    # --- Indeks ---
    lines = ["# Indeks kebutuhan (baca ini dulu)\n",
             "Sumber kebenaran: BRD v4.0 dan FRD v3.0 (xlsx di root repo). File di sini dibuat otomatis; ",
             "perubahan isi kebutuhan harus disetujui Gab, lalu xlsx diperbarui dan script dijalankan ulang.\n",
             "## Modul FRD (baca hanya modul yang dikerjakan)\n",
             "| Kode | Modul | Kebutuhan | Must | File |", "| --- | --- | --- | --- | --- |"]
    for code, name, n, must in index_mod:
        lines.append(f"| {code} | {cell(name)} | {n} | {must} | frd/modules/{code}.md |")
    lines.append("\n## File lain\n")
    lines.append("| File | Baris |")
    lines.append("| --- | --- |")
    for p in written:
        n = p.read_text(encoding="utf-8").count("\n")
        lines.append(f"| {p.relative_to(ROOT).as_posix()} | {n} |")
    write(OUT_REQ / "INDEX.md", "\n".join(lines) + "\n", "BRD/FRD")
    print(f"{len(written)} file + INDEX.md")


if __name__ == "__main__":
    sys.exit(main())
